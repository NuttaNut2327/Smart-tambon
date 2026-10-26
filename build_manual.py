from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(r"C:\Users\nutta\.codex\attachments\8eb10c5e-14d3-45c1-9b29-c53967c9393c\Pasted text.txt")
OUT = Path(r"F:\ESM_work\smart-tambon\คู่มือการใช้งาน Smart Tambon ฉบับปรับปรุง.docx")

lines = [x.strip() for x in SRC.read_text(encoding="utf-8-sig").splitlines()]

markers = {
    "start": "1.\tเริ่มต้นใช้งาน",
    "dashboard": "2.\tภาพรวมเมือง",
    "map": "3.\tแผนที่เมือง",
    "incident": "4.\tเหตุการณ์ ภัยพิบัติ และแจ้งเหตุจากประชาชน",
    "rules": "5. กฎการประเมินทรัพยากร",
    "activity": "6. กิจกรรมในพื้นที่",
    "news": "7. ข่าวประชาสัมพันธ์ เสียงตามสาย และช่องทางแจ้งเตือน",
    "cctv": "8. กล้องวงจรปิด",
    "project": "9. โครงการ",
    "import": "10. นำเข้าข้อมูล",
    "ai": "11. ผู้ช่วย AI",
    "admin": "8. งานผู้ดูแลระบบ",
    "faq": "12. คำถามที่พบบ่อย",
}
pos = {k: lines.index(v) for k, v in markers.items()}
ordered_keys = ["start", "dashboard", "map", "incident", "rules", "activity", "news", "cctv", "project", "import", "ai", "admin", "faq"]
blocks = {}
for i, key in enumerate(ordered_keys):
    a = pos[key] + 1
    b = pos[ordered_keys[i + 1]] if i + 1 < len(ordered_keys) else len(lines)
    blocks[key] = [x for x in lines[a:b] if x]

chapters = [
    ("1 เริ่มต้นใช้งาน", [("start", None)]),
    ("2 ภาพรวมเมืองและแผนที่", [("dashboard", "2.1 ภาพรวมเมือง"), ("map", "2.2 แผนที่เมือง")]),
    ("3 การนำเข้าข้อมูล", [("import", None)]),
    ("4 เหตุการณ์ ภัยพิบัติ และการแจ้งเหตุจากประชาชน", [("incident", "4.1 การจัดการเหตุการณ์และภัยพิบัติ"), ("rules", "4.2 กฎการประเมินทรัพยากร")]),
    ("5 กิจกรรม โครงการ กล้องวงจรปิด และข้อมูลเฝ้าระวัง", [("activity", "5.1 กิจกรรมในพื้นที่"), ("project", "5.2 โครงการ"), ("cctv", "5.3 กล้องวงจรปิดและข้อมูลเฝ้าระวัง")]),
    ("6 ผู้ช่วย AI", [("ai", None)]),
    ("7 ข่าวประชาสัมพันธ์ เสียงตามสาย และช่องทางแจ้งเตือน", [("news", None)]),
    ("8 งานสำหรับผู้ดูแลระบบ", [("admin", None)]),
    ("9 คำถามที่พบบ่อย", [("faq", None)]),
]

replacements = {
    "คู่มือนี้อธิบายทุกเมนูที่ใช้งานได้ในระบบปัจจุบัน ชื่อปุ่มและเมนูในเครื่องหมาย “…” ตรงกับที่เห็นบนหน้าจอ":
        "คู่มือนี้อธิบายวิธีใช้งานเมนูที่มีอยู่ในระบบปัจจุบัน โดยชื่อปุ่มและชื่อเมนูที่อยู่ในเครื่องหมายอัญประกาศเป็นชื่อเดียวกับที่ปรากฏบนหน้าจอ",
    "ถามเป็นภาษาไทยได้เลย ระบบจะค้นข้อมูลในพื้นที่ของคุณแล้วตอบเป็นข้อความ กราฟ หรือตาราง":
        "ผู้ใช้งานสามารถสอบถามเป็นภาษาไทยได้ ระบบจะค้นหาข้อมูลภายในพื้นที่ที่บัญชีมีสิทธิ์เข้าถึง และแสดงคำตอบในรูปแบบข้อความ กราฟ หรือตาราง",
    "เจ้าหน้าที่ทั่วไปดูได้อย่างเดียว ผู้ดูแลเป็นผู้เพิ่ม/แก้ไข":
        "เจ้าหน้าที่ประจำ อบต. สามารถดูข้อมูลได้เท่านั้น ส่วนการเพิ่มหรือแก้ไขข้อมูลเป็นสิทธิ์ของผู้ดูแล",
    "ไม่ต้องล็อกอิน": "ไม่ต้องเข้าสู่ระบบ",
    "ถูกดีดออกทันที": "ระบบนำกลับไปยังหน้าเข้าสู่ระบบทันที",
    "ติ๊ก ": "เลือกช่อง ",
    "แล้ว เสร็จสิ้น": "แล้วเลือก “เสร็จสิ้น”",
    "ใส่ “รูปหลัก": "ระบุ “รูปหลัก",
    "ใส่วันเวลา": "ระบุวันและเวลา",
    "ใส่โมเดล": "ระบุโมเดล",
    "เบอร์ (แสดง": "หมายเลขโทรศัพท์ (แสดง",
    "เบอร์ติดต่อ": "หมายเลขโทรศัพท์ติดต่อ",
    "เบอร์ผู้แจ้งเหตุ": "หมายเลขโทรศัพท์ของผู้แจ้งเหตุ",
    "ทุกหน้าอื่น": "หน้าส่วนอื่นทั้งหมด",
    "ไฟล์ที่ค้างกลางทาง": "ไฟล์ที่อยู่ระหว่างขั้นตอนนำเข้า",
    "ทิ้งไว้นาน": "ไม่มีการดำเนินการภายในระยะเวลาดังกล่าว",
    "กดปุ่ม": "เลือกปุ่ม",
    "กด “": "เลือก “",
}

def formalize(text):
    text = text.replace("—", "-").replace("→", "ไปยัง")
    for a, b in replacements.items():
        text = text.replace(a, b)
    text = re.sub(r"\s+", " ", text).strip()
    return text

def set_font(run, name="TH Sarabun New", size=16, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table):
    tblPr = table._tbl.tblPr
    element = tblPr.first_child_found_in("w:tblBorders")
    if element is None:
        element = OxmlElement("w:tblBorders")
        tblPr.append(element)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")
        element.append(tag)

def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.0)

styles = doc.styles
for style_name, size, bold, before, after in [
    ("Normal", 16, False, 0, 5), ("Title", 28, True, 0, 10),
    ("Heading 1", 22, True, 12, 8), ("Heading 2", 19, True, 9, 5),
    ("Heading 3", 17, True, 6, 3),
]:
    st = styles[style_name]
    st.font.name = "TH Sarabun New"; st.font.size = Pt(size); st.font.bold = bold; st.font.color.rgb = RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing = 1.05

# Cover
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("คู่มือการใช้งาน Smart Tambon")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ระบบบริหารจัดการความปลอดภัยอัจฉริยะสำหรับองค์การบริหารส่วนตำบล"); set_font(r, size=20, bold=True)
doc.add_paragraph("")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("จัดทำโดย E.S.M. Solution Co., Ltd."); set_font(r, size=17)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ปรับปรุงล่าสุด 22 ตุลาคม 2569"); set_font(r, size=16)
doc.add_page_break()

doc.add_heading("เกี่ยวกับคู่มือ", level=1)
p = doc.add_paragraph("คู่มือนี้จัดทำขึ้นเพื่ออธิบายขั้นตอนการใช้งานระบบ Smart Tambon สำหรับผู้ดูแลระบบ ผู้ดูแลประจำองค์การบริหารส่วนตำบล และเจ้าหน้าที่ประจำองค์การบริหารส่วนตำบล เนื้อหาครอบคลุมการตั้งค่าเริ่มต้น การบริหารข้อมูลพื้นที่ การรับและติดตามเหตุการณ์ การประเมินภัยพิบัติ การจัดการกิจกรรมและโครงการ ตลอดจนการใช้ผู้ช่วย AI และการเผยแพร่ข่าวประชาสัมพันธ์")
p.paragraph_format.first_line_indent = Cm(0.75)
doc.add_heading("สารบัญ", level=1)
toc = doc.add_table(rows=1, cols=2)
toc.alignment = WD_TABLE_ALIGNMENT.CENTER
toc.columns[0].width, toc.columns[1].width = Cm(1.2), Cm(15.5)
toc.rows[0].cells[0].text, toc.rows[0].cells[1].text = "บท", "หัวข้อ"
for idx, (title, _) in enumerate(chapters, 1):
    row = toc.add_row().cells
    row[0].text = str(idx); row[1].text = re.sub(r"^\d+\s+", "", title)
for i, row in enumerate(toc.rows):
    for j, cell in enumerate(row.cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margin(cell)
        shade(cell, "1F4E78" if i == 0 else ("EDF3F8" if i % 2 == 0 else "FFFFFF"))
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs: set_font(r, size=15, bold=(i == 0), color="FFFFFF" if i == 0 else "000000")
borders(toc)

known_subheads = {
    "เข้าสู่ระบบ", "เมนูด้านซ้าย", "สิทธิ์ผู้ใช้งาน", "3.1. วิเคราะห์พื้นที่", "3.2. จำลองสถานการณ์",
    "4.1. เจ้าหน้าที่บันทึกเหตุการณ์", "4.2. แบบฟอร์มแจ้งเหตุจากประชาชน", "1. รับเรื่องและติดตาม",
    "4.4. ภัยพิบัติและแผนรับมือ (ผู้ดูแล)", "ตั้ง “รายการข่าว” (เช่น ข่าวเช้า 07:00)",
    "ตรวจแก้และเผยแพร่ข่าว", "ข้อมูลเฝ้าระวัง (อัตโนมัติ)", "ประเภทข้อมูล", "นำเข้าไฟล์ (6 ขั้น)",
    "เพิ่ม/แก้ไขทีละรายการ", "รุ่นข้อมูลและประวัติ", "ถามคำถาม", "ตัวอย่างคำถาม", "เก็บรายงานไว้ใช้ซ้ำ",
    "ช่องทางแจ้งเตือน (ผู้ดูแลประจำ อบต.)", "จัดการผู้ใช้", "ตั้งค่า อบต. (ผู้ดูแลประจำ อบต.)",
    "ตั้งค่าทั้งระบบ (ผู้ดูแลระบบ)",
}
subhead_renames = {
    "3.1. วิเคราะห์พื้นที่": "การวิเคราะห์พื้นที่", "3.2. จำลองสถานการณ์": "การจำลองสถานการณ์",
    "4.1. เจ้าหน้าที่บันทึกเหตุการณ์": "การบันทึกเหตุการณ์โดยเจ้าหน้าที่",
    "4.2. แบบฟอร์มแจ้งเหตุจากประชาชน": "แบบฟอร์มแจ้งเหตุสำหรับประชาชน",
    "1. รับเรื่องและติดตาม": "การรับเรื่องและติดตามผล",
    "4.4. ภัยพิบัติและแผนรับมือ (ผู้ดูแล)": "การจัดการภัยพิบัติและแผนรับมือสำหรับผู้ดูแล",
    "ตั้ง “รายการข่าว” (เช่น ข่าวเช้า 07:00)": "การตั้งค่ารายการข่าว",
    "ตรวจแก้และเผยแพร่ข่าว": "การตรวจสอบและเผยแพร่ข่าว",
    "ข้อมูลเฝ้าระวัง (อัตโนมัติ)": "ข้อมูลเฝ้าระวังอัตโนมัติ",
    "นำเข้าไฟล์ (6 ขั้น)": "ขั้นตอนการนำเข้าไฟล์",
    "เพิ่ม/แก้ไขทีละรายการ": "การเพิ่มและแก้ไขข้อมูลรายรายการ",
    "ตั้งค่า อบต. (ผู้ดูแลประจำ อบต.)": "การตั้งค่าองค์การบริหารส่วนตำบล",
    "ตั้งค่าทั้งระบบ (ผู้ดูแลระบบ)": "การตั้งค่าระบบส่วนกลาง",
}

def is_table_start(items, i):
    return "\t" in items[i] and not re.match(r"^\d+[.)]\s", items[i])

def add_table_rows(rows):
    maxcols = max(len(r.split("\t")) for r in rows)
    table = doc.add_table(rows=0, cols=maxcols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, raw in enumerate(rows):
        vals = [formalize(x) for x in raw.split("\t")]
        vals += [""] * (maxcols - len(vals))
        cells = table.add_row().cells
        for ci, (cell, val) in enumerate(zip(cells, vals)):
            cell.text = val; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margin(cell)
            shade(cell, "1F4E78" if ri == 0 else ("EDF3F8" if ri % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci == 0 or (maxcols > 2 and ci > 1) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs: set_font(r, size=14, bold=(ri == 0), color="FFFFFF" if ri == 0 else "000000")
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

for ci, (chapter, parts) in enumerate(chapters):
    doc.add_page_break()
    doc.add_heading(chapter, level=1)
    for key, section_title in parts:
        if section_title:
            doc.add_heading(section_title, level=2)
        items = blocks[key]
        i = 0
        while i < len(items):
            line = items[i]
            if is_table_start(items, i):
                rows = []
                while i < len(items) and "\t" in items[i]:
                    rows.append(items[i]); i += 1
                if len(rows) >= 2:
                    add_table_rows(rows); continue
            if line in known_subheads:
                doc.add_heading(subhead_renames.get(line, re.sub(r"^\d+(?:\.\d+)?[.)]?\s*", "", line)), level=3)
            elif line.startswith("หน้") and ":" in line:
                p = doc.add_paragraph()
                r = p.add_run(formalize(line)); set_font(r, size=14, bold=False, color="666666"); r.italic = True
                p.paragraph_format.keep_with_next = True
            elif line.startswith("หมายเหตุ:"):
                p = doc.add_paragraph()
                r = p.add_run("หมายเหตุ "); set_font(r, bold=True)
                r = p.add_run(formalize(line.split(":",1)[1])); set_font(r)
                p.paragraph_format.left_indent = Cm(0.75)
            elif re.match(r"^[•]\s*", line):
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(formalize(re.sub(r"^[•]\s*", "", line)))
            elif re.match(r"^\d+[.)]\s*", line):
                p = doc.add_paragraph(style="List Number")
                p.add_run(formalize(re.sub(r"^\d+[.)]\s*", "", line)))
            else:
                p = doc.add_paragraph(formalize(line))
                p.paragraph_format.first_line_indent = Cm(0.75)
            i += 1

for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("คู่มือการใช้งาน Smart Tambon")
    set_font(r, size=11, color="666666")
    r = p.add_run("   |   ")
    set_font(r, size=11, color="999999")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)

doc.core_properties.title = "คู่มือการใช้งาน Smart Tambon"
doc.core_properties.subject = "ระบบบริหารจัดการความปลอดภัยอัจฉริยะสำหรับองค์การบริหารส่วนตำบล"
doc.core_properties.author = "E.S.M. Solution Co., Ltd."
doc.save(OUT)
print(OUT)
